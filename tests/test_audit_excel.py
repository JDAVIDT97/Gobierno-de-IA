from datetime import datetime
from io import BytesIO
from types import SimpleNamespace

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from api.audit_excel import HEADERS, PROVIDERS, build_audit_workbook, provider_name
from api.database import AuditLog, Base, get_db
from api.main import app


def record(ident=1, destination="https://chatgpt.com/", **overrides):
    values = dict(id=ident, timestamp="2026-10-31T02:30:00+00:00", user_id="synthetic-user",
                  destination=destination, action="BLOCK", category="PII_Leak", risk_score=0.85,
                  inference_time_ms=4.125)
    values.update(overrides)
    return SimpleNamespace(**values)


@pytest.mark.parametrize("destination,expected", [
    ("https://chatgpt.com/uc/test", "ChatGPT"),
    ("https://chat.openai.com/c/test", "ChatGPT"),
    ("https://claude.ai/new", "Claude"),
    ("https://copilot.cloud.microsoft/", "Copilot"),
    ("https://m365.cloud.microsoft/chat", "Copilot"),
    ("https://copilot.com/", "Copilot"),
    ("https://gemini.google.com/app", "Gemini"),
    ("https://x.com/i/grok", "Grok"),
    ("https://chat.deepseek.com/", "DeepSeek"),
    ("https://chatgpt.com.evil.invalid/", "Otros destinos"),
    ("http://127.0.0.1:8000/sandbox", "Pruebas locales"),
    (None, "Sin destino"),
])
def test_provider_uses_destination_host(destination, expected):
    assert provider_name(destination) == expected


def test_workbook_preserves_values_types_and_separates_providers():
    workbook = load_workbook(build_audit_workbook([
        record(), record(2, "https://claude.ai/", user_id="=1+1"), record(3, None),
    ]))
    assert list(workbook.sheetnames[:len(PROVIDERS)]) == list(PROVIDERS)
    chatgpt = workbook["ChatGPT"]
    assert [chatgpt.cell(5, col).value for col in range(1, 9)] == HEADERS
    assert chatgpt["B6"].value == datetime(2026, 10, 30, 21, 30)
    assert chatgpt["G6"].value == 0.85
    assert chatgpt["G6"].number_format == "0.0%"
    assert chatgpt["H6"].value == 4.125
    assert chatgpt.freeze_panes == "A6"
    assert len(chatgpt.tables) == 1
    assert workbook["Claude"]["A6"].value == 2
    assert workbook["Claude"]["C6"].value == "=1+1"
    assert workbook["Claude"]["C6"].data_type == "s"
    assert workbook["Sin destino"]["A6"].value == 3
    assert workbook["Gemini"]["B2"].value == 0


def test_empty_history_creates_valid_empty_provider_tables():
    workbook = load_workbook(build_audit_workbook([]))
    assert workbook.sheetnames == list(PROVIDERS)
    assert all(sheet["B2"].value == 0 and sheet["A6"].value is None for sheet in workbook)
    assert all(len(sheet.tables) == 1 for sheet in workbook)


def test_excel_endpoint_exports_entire_history_without_pagination():
    engine = create_engine("sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    session_factory = sessionmaker(bind=engine)
    with session_factory() as db:
        for ident in range(1, 252):
            row = vars(record(ident, "https://chatgpt.com" if ident <= 250 else "https://claude.ai"))
            db.add(AuditLog(**row, prompt_hash="synthetic", prompt_redacted="[SYNTHETIC]",
                            confidence=1, model_version="test"))
        db.commit()

    def override_db():
        with session_factory() as db:
            yield db

    original = app.dependency_overrides.copy()
    app.dependency_overrides[get_db] = override_db
    try:
        response = TestClient(app).get("/v1/export-excel")
        assert response.status_code == 200
        assert "spreadsheetml.sheet" in response.headers["content-type"]
        assert response.headers["content-disposition"].endswith("aegis_historial_por_IA.xlsx")
        workbook = load_workbook(BytesIO(response.content))
        assert workbook["ChatGPT"]["B2"].value == 250
        assert workbook["ChatGPT"]["A6"].value == 250
        assert workbook["ChatGPT"]["A255"].value == 1
        assert workbook["Claude"]["B2"].value == 1
        assert sum(sheet["B2"].value for sheet in workbook) == 251
    finally:
        app.dependency_overrides.clear()
        app.dependency_overrides.update(original)
        engine.dispose()
