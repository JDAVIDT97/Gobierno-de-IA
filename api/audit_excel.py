"""Excel export of the complete audit history, separated by AI destination."""
from collections import OrderedDict
from datetime import datetime, timezone
from io import BytesIO
from urllib.parse import urlsplit
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo


PERU = ZoneInfo("America/Lima")
PROVIDERS = OrderedDict([
    ("ChatGPT", ("chatgpt.com", "chat.openai.com")),
    ("Claude", ("claude.ai",)),
    ("Gemini", ("gemini.google.com",)),
    ("Copilot", ("copilot.microsoft.com", "copilot.com", "copilot.cloud.microsoft", "m365.cloud.microsoft")),
    ("Perplexity", ("perplexity.ai", "www.perplexity.ai")),
    ("Grok", ("grok.com", "x.com")),
    ("Mistral", ("chat.mistral.ai",)),
    ("DeepSeek", ("chat.deepseek.com",)),
    ("Poe", ("poe.com",)),
    ("You.com", ("you.com",)),
])
HEADERS = ["ID", "Fecha y hora (Perú)", "Usuario", "Destino", "Decisión", "Categoría", "Riesgo", "Latencia"]
MAX_DATA_ROWS = 1_048_571  # Excel's row limit, reserving five heading rows.


def provider_name(destination):
    if not destination:
        return "Sin destino"
    try:
        host = (urlsplit(destination).hostname or "").lower()
    except ValueError:
        return "Otros destinos"
    for name, hosts in PROVIDERS.items():
        if any(host == allowed or host.endswith("." + allowed) for allowed in hosts):
            return name
    if host in {"localhost", "127.0.0.1", "::1"}:
        return "Pruebas locales"
    return "Otros destinos"


def peru_datetime(value):
    try:
        parsed = datetime.fromisoformat(value.replace("Z", "+00:00"))
        if parsed.tzinfo is None:
            parsed = parsed.replace(tzinfo=timezone.utc)
        return parsed.astimezone(PERU).replace(tzinfo=None)
    except (ValueError, AttributeError, TypeError):
        return value or ""


def build_audit_workbook(records):
    groups = OrderedDict((name, []) for name in PROVIDERS)
    for record in records:
        groups.setdefault(provider_name(record.destination), []).append(record)
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.title = "Historial de auditoría por IA"
    workbook.properties.creator = "Aegis DLP"
    generated = datetime.now(PERU).replace(tzinfo=None)
    for provider, rows in groups.items():
        for start in range(0, max(len(rows), 1), MAX_DATA_ROWS):
            part = start // MAX_DATA_ROWS + 1
            title = provider if len(rows) <= MAX_DATA_ROWS else f"{provider} {part}"
            sheet = workbook.create_sheet(title)
            sheet.sheet_view.showGridLines = False
            sheet.merge_cells("A1:H1")
            sheet["A1"] = provider + " — Bitácora de auditoría"
            sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color="FFFFFF")
            sheet["A1"].fill = PatternFill("solid", fgColor="10243A")
            sheet.row_dimensions[1].height = 32
            chunk = rows[start:start + MAX_DATA_ROWS]
            sheet["A2"] = "Registros"
            sheet["B2"] = len(chunk)
            sheet["C2"] = "Exportado (Perú)"
            sheet["D2"] = generated
            sheet["D2"].number_format = "dd/mm/yyyy hh:mm:ss"
            sheet["A3"] = "Riesgo en porcentaje. Latencia en milisegundos. Historial completo guardado."
            sheet.merge_cells("A3:H3")
            sheet["A3"].font = Font(name="Calibri", size=10, color="526174")
            for column, heading in enumerate(HEADERS, 1):
                cell = sheet.cell(5, column, heading)
                cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="176B87")
                cell.alignment = Alignment(vertical="center")
            sheet.row_dimensions[5].height = 26
            for row_number, record in enumerate(chunk, 6):
                values = [record.id, peru_datetime(record.timestamp), record.user_id,
                          record.destination or "", record.action, record.category,
                          record.risk_score, record.inference_time_ms]
                for column, value in enumerate(values, 1):
                    cell = sheet.cell(row_number, column, value)
                    # Treat source text literally, including strings beginning with '='.
                    if isinstance(value, str):
                        cell.data_type = "s"
                    cell.font = Font(name="Calibri", size=11)
                    cell.alignment = Alignment(vertical="center")
                sheet.cell(row_number, 2).number_format = "dd/mm/yyyy hh:mm:ss"
                sheet.cell(row_number, 7).number_format = "0.0%"
                sheet.cell(row_number, 8).number_format = '0.00" ms"'
                decision = sheet.cell(row_number, 5)
                decision.font = Font(name="Calibri", bold=True,
                                     color="B91C1C" if record.action == "BLOCK" else "166534")
            table = Table(displayName=f"Auditoria_{len(workbook.worksheets)}",
                          ref=f"A5:H{max(6, 5 + len(chunk))}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
            for column, width in zip("ABCDEFGH", (10, 24, 40, 48, 16, 30, 14, 18)):
                sheet.column_dimensions[column].width = width
            sheet.freeze_panes = "A6"
            sheet.sheet_properties.pageSetUpPr.fitToPage = True
            sheet.page_setup.orientation = "landscape"
            sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
            sheet.page_setup.fitToWidth = 1
            sheet.page_setup.fitToHeight = 0
            sheet.print_title_rows = "1:5"
            sheet.print_options.horizontalCentered = True
    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output
